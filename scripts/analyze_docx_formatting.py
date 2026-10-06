import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

def analyze_docx_formatting(path):
    print(f"\n==========================================")
    print(f"File: {path}")
    print(f"==========================================")
    doc = Document(path)
    
    total_p = len(doc.paragraphs)
    print(f"Total paragraphs: {total_p}")
    
    bold_count = 0
    highlight_count = 0
    color_count = 0
    underline_count = 0
    
    for idx, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t: continue
        
        has_b = any(r.bold for r in p.runs if r.text.strip())
        has_h = any(r.font.highlight_color for r in p.runs if r.text.strip())
        has_c = any(r.font.color and r.font.color.rgb for r in p.runs if r.text.strip())
        has_u = any(r.underline for r in p.runs if r.text.strip())
        
        if has_b: bold_count += 1
        if has_h: highlight_count += 1
        if has_c: color_count += 1
        if has_u: underline_count += 1
        
        if has_b or has_h or has_c or has_u:
            if idx < 60:
                flags = []
                if has_b: flags.append("BOLD")
                if has_h: flags.append("HIGHLIGHT")
                if has_c: flags.append("COLOR")
                if has_u: flags.append("UNDERLINE")
                print(f"P{idx+1} [{' '.join(flags)}]: {t}")
                
    print(f"Summary: bold_p={bold_count}, highlight_p={highlight_count}, color_p={color_count}, underline_p={underline_count}")

analyze_docx_formatting('raw_materials/exams/שחזור הדרכה קלינית 23-12-25.docx')
analyze_docx_formatting('raw_materials/exams/שחזור מבחן מאי הדרכה קלינית.docx')

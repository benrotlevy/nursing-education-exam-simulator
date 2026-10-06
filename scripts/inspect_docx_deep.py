import os
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

docx_files = [
    'raw_materials/exams/שחזור מבחן אמצע הדרכה קלינית 06.docx',
    'raw_materials/exams/שחזור הדרכה קלינית 23-12-25.docx',
    'raw_materials/exams/שחזור מבחן מאי הדרכה קלינית.docx'
]

for path in docx_files:
    print(f"\n==========================================")
    print(f"Inspecting: {os.path.basename(path)}")
    print(f"==========================================")
    doc = Document(path)
    
    print(f"Paragraphs count: {len(doc.paragraphs)}")
    print(f"Tables count: {len(doc.tables)}")
    
    print("--- First 30 paragraphs with formatting ---")
    for i, p in enumerate(doc.paragraphs[:30]):
        text = p.text.strip()
        if not text:
            continue
        # Check if any run in paragraph is bold or colored or highlighted
        bolds = [r.text for r in p.runs if r.bold and r.text.strip()]
        highs = [r.text for r in p.runs if r.font.highlight_color and r.text.strip()]
        colors = [r.text for r in p.runs if r.font.color and r.font.color.rgb and r.text.strip()]
        
        format_info = []
        if bolds: format_info.append(f"BOLD: {' '.join(bolds)}")
        if highs: format_info.append(f"HIGHLIGHT: {' '.join(highs)}")
        if colors: format_info.append(f"COLOR: {' '.join(colors)}")
        
        f_str = f" [{', '.join(format_info)}]" if format_info else ""
        print(f"P{i+1}: {text}{f_str}")
        
    if doc.tables:
        print(f"--- Table 1 preview ({len(doc.tables[0].rows)} rows) ---")
        for r in doc.tables[0].rows[:5]:
            cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(" | ".join(cells))
